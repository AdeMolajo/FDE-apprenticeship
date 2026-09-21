"""Build a deliberately messy data-room spreadsheet (xlsx) plus a CSV export
with genuine byte-level encoding problems."""

import csv
import io
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

out_dir = Path(sys.argv[1])
out_dir.mkdir(parents=True, exist_ok=True)

NBSP = " "


def d(y, m, day):  # a real Excel date cell
    return datetime(y, m, day)


# --- Sheet 1: FY25 sales ledger, two exports pasted together -----------------
header_a = ["Invoice No", "Customer Name ", "invoice_date", "Amount (GBP)", "Region", "Status"]
rows_a = [
    ["INV-2501", "Acme Motors Ltd", "2025-01-14", 12500.00, "North", "Paid"],
    ["INV-2502", "SociÃ©tÃ© GÃ©nÃ©rale UK", "14/02/2025", "1,250.00", "London", "Paid"],
    ["INV-2503", "MÃ¼ller Engineering GmbH", d(2025, 3, 3), "Â£3,400", "EU", "Outstanding"],
    ["INV-2504", "Caf� Nero Holdings", "2025-03-19", 870, "South", "Paid â€“ partial"],
    ["INV-2505", "Northwind Traders", "03/04/2025", "(500.00)", "North", "Credit note"],
    ["INV-2506", "Blue Harbour Logistics", d(2025, 4, 22), "N/A", "Scotland", "Disputed"],
    ["INV-2507", "Acme Motors Ltd", "2025-05-02", f"2{NBSP}100.50", "North", "Paid"],
    ["INV-2507", "Acme Motors Ltd", "2025-05-02", f"2{NBSP}100.50", "North", "Paid"],  # exact dup
    ["INV-2508", "Ãrsted Wind Services", "11/05/2025", "TBC", "EU", "Outstanding"],
    ["INV-2509", "Greenfield Farms", d(2025, 6, 1), 4300, "Wales", "paid"],
    ["INV-2510", "Greenfield Farms ", d(2025, 6, 1), 4300, "Wales", "Paid"],  # near dup: space, case
]
header_b = ["Inv #", "Client", "Date of Invoice", "Amount", "region", "STATUS"]  # second export
rows_b = [
    ["INV-2511", "Blue Harbour Logistics", "2025-07-08", "6,780.00", "scotland", "Paid"],
    ["INV-2512", "ZÃ¼rich Insurance plc", "08/07/2025", 1.2e3, "London", "Outstanding"],
    ["INV-2513", "Northwind Traders", d(2025, 7, 30), "-", "North", ""],
    ["INV-2514", "Acme Motors Ltd", "2025-08-15", "15000", "North", "Paid"],
    ["INV-2514", "Acme Motors Ltd", "2025-08-15", "15000", "North", "Paid"],  # exact dup
    ["INV-2515", "O’Connor & Sons", "01/09/2025", "Â£ 950.00", "NI", "Paid"],
    ["INV-2516", "SmërgÃ¥sbord Foods AB", d(2025, 9, 12), True, "EU", "Outstanding"],
    ["", "", "", "", "", ""],
    ["TOTAL", "", "", "=SUM(D2:D24)", "", ""],
]

wb = Workbook()
ws = wb.active
ws.title = "Sales FY25"
for row in [header_a, *rows_a, header_b, *rows_b]:
    ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_col=3, max_col=3):
    if isinstance(row[0].value, datetime):
        row[0].number_format = "d-mmm-yy"

# --- Sheet 2: FY24, same data model, different headers -----------------------
ws2 = wb.create_sheet("sales fy24 ")
ws2.append(["InvoiceNumber", "Customer", "Date", "Value £", "Territory", "Paid?"])
for row in [
    ["INV-2401", "Acme Motors Ltd", "2024-02-10", 11200, "North", "Y"],
    ["INV-2402", "MÃ¼ller Engineering GmbH", "10/03/2024", "2.950,00", "EU", "yes"],
    ["INV-2403", "Northwind Traders", d(2024, 4, 5), "3100", "North", 1],
    ["INV-2404", "Blue Harbour Logistics", "2024-06-30", "£4,050", "Scotland", "N"],
    ["INV-2404", "Blue Harbour Logistics", "2024-06-30", "£4,050", "Scotland", "N"],  # exact dup
    ["INV-2405", "Greenfield Farms", "07/08/2024", "n/a", "Wales", "No"],
    ["INV-2406", "SociÃ©tÃ© GÃ©nÃ©rale UK", d(2024, 11, 18), 7625.5, "London", True],
]:
    ws2.append(row)
for row in ws2.iter_rows(min_col=3, max_col=3):
    if isinstance(row[0].value, datetime):
        row[0].number_format = "d-mmm-yy"

xlsx_path = out_dir / "Sales Ledger FY24-FY25 (FINAL v3).xlsx"
wb.save(xlsx_path)

# --- CSV export of FY25 with real encoding problems ---------------------------
# Mostly UTF-8 with a BOM, but a few rows were pasted in from a Windows-1252 file,
# so their accented characters are raw cp1252 bytes (invalid as UTF-8).
cp1252_rows = {"INV-2502", "INV-2503", "INV-2516"}
clean_names = {
    "INV-2502": "Société Générale UK",
    "INV-2503": "Müller Engineering GmbH",
    "INV-2516": "Smörgåsbord Foods AB",
}
buf = io.BytesIO()
buf.write("﻿".encode("utf-8"))  # BOM
for row in [header_a, *rows_a, header_b, *rows_b[:-2]]:
    row = ["" if v is None else v for v in row]
    row = [v.strftime("%d-%b-%y") if isinstance(v, datetime) else v for v in row]
    inv = row[0]
    if inv in cp1252_rows:
        row[1] = clean_names[inv]
        line = io.StringIO()
        csv.writer(line, lineterminator="\r\n").writerow(row)
        buf.write(line.getvalue().encode("cp1252", errors="replace"))
    else:
        line = io.StringIO()
        csv.writer(line, lineterminator="\n").writerow(row)  # mixed line endings too
        buf.write(line.getvalue().encode("utf-8"))
csv_path = out_dir / "Sales Ledger FY25 export.csv"
csv_path.write_bytes(buf.getvalue())

print(xlsx_path)
print(csv_path)
